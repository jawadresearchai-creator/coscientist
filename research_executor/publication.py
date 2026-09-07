from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import tempfile
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Callable

from research_executor.acquisition import (
    AcquisitionError,
    _child,
    _download_json,
    _drive_client_from_env,
    _optional_child,
    _search_job_folder,
    _write_json,
    sha256_file,
)
from research_executor.public_status import validate_job_id, write_public_status

REGISTRY_PATH = Path(__file__).with_name("publication_registry.json")
_DOI_RE = re.compile(r"^10\.\d{4,9}/\S+$", re.I)
_SAFE_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")


class PublicationError(AcquisitionError):
    pass


def canonical_sha256(value: Any) -> str:
    raw = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def load_publication_registry(path: str | Path = REGISTRY_PATH) -> dict[str, Any]:
    registry = json.loads(Path(path).read_text(encoding="utf-8"))
    if registry.get("schema_version") != "cosci.publication-registry/1.0":
        raise PublicationError("unexpected publication registry schema")
    policy = registry.get("policy") or {}
    required_true = [
        "zero_cost_only", "private_manuscript_only", "official_journal_requirements_required",
        "journal_requirements_provenance_required", "single_canonical_source_required",
        "artifact_hash_validation_required", "citation_verification_required",
        "figure_table_traceability_required", "docx_required", "pdf_required",
        "deterministic_build_required", "checkpoint_required", "idempotent_rerun_required",
        "journal_specific_overrides_take_precedence", "acceptance_fixture_must_be_labeled",
        "final_scientific_judgment_reserved_for_research_director",
    ]
    for key in required_true:
        if policy.get(key) is not True:
            raise PublicationError(f"publication policy must enforce {key}=true")
    if policy.get("automatic_paid_routes") is not False:
        raise PublicationError("publication policy must disable automatic paid routes")
    return registry


def normalize_doi(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    for prefix in ("https://doi.org/", "http://doi.org/", "doi:"):
        if text.startswith(prefix):
            text = text[len(prefix):]
    text = text.rstrip(". ")
    return text if _DOI_RE.fullmatch(text) else None


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignore = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self._ignore += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self._ignore:
            self._ignore -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignore:
            cleaned = " ".join(html.unescape(data).split())
            if cleaned:
                self.parts.append(cleaned)


def _request_bytes(
    url: str,
    *,
    opener: Callable[..., Any] = urllib.request.urlopen,
    max_bytes: int = 2_000_000,
) -> tuple[bytes, dict[str, str]]:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Research-CoScientist-Publication/1.0",
            "Accept": "text/html,application/json;q=0.9,*/*;q=0.8",
        },
    )
    try:
        with opener(request, timeout=120) as response:
            raw = response.read(max_bytes + 1)
            headers = {
                str(k).lower(): str(v)
                for k, v in getattr(response, "headers", {}).items()
            }
    except Exception as exc:
        raise PublicationError(
            f"official zero-cost source request failed: {type(exc).__name__}"
        ) from exc
    if len(raw) > max_bytes:
        raise PublicationError("official source response exceeds publication safety limit")
    return raw, headers


def retrieve_journal_requirements(
    spec: dict[str, Any],
    *,
    opener: Callable[..., Any] = urllib.request.urlopen,
) -> dict[str, Any]:
    url = str(spec.get("official_url") or "")
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme != "https" or not parsed.netloc:
        raise PublicationError("journal requirements must use an official HTTPS route")
    raw, headers = _request_bytes(url, opener=opener)
    content_type = headers.get("content-type", "")
    if "json" in content_type:
        try:
            obj = json.loads(raw.decode("utf-8"))
            text = json.dumps(obj, ensure_ascii=False, sort_keys=True)
        except Exception as exc:
            raise PublicationError("official journal requirement JSON was invalid") from exc
    else:
        decoded = raw.decode("utf-8", errors="replace")
        parser = _TextExtractor()
        parser.feed(decoded)
        text = "\n".join(parser.parts)
    markers = [str(x) for x in (spec.get("required_markers") or []) if str(x).strip()]
    lowered = text.lower()
    missing = [m for m in markers if m.lower() not in lowered]
    if missing:
        raise PublicationError(
            "official journal requirement response failed expected-marker validation"
        )
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return {
        "schema_version": "cosci.journal-requirements/1.0",
        "journal_key": str(spec.get("journal_key") or "PRIVATE_TARGET_JOURNAL"),
        "source_type": "OFFICIAL_HTTPS",
        "official_url": url,
        "content_sha256": hashlib.sha256(raw).hexdigest(),
        "normalized_text_sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "required_markers": markers,
        "marker_validation": "PASS",
        "content_type": content_type.split(";", 1)[0],
        "normalized_text": normalized[:250_000],
    }


def verify_doi(
    doi: str,
    *,
    opener: Callable[..., Any] = urllib.request.urlopen,
) -> dict[str, Any]:
    norm = normalize_doi(doi)
    if not norm:
        return {"doi": str(doi), "status": "INVALID_FORMAT"}
    encoded = urllib.parse.quote(norm, safe="")
    raw, _ = _request_bytes(
        f"https://api.crossref.org/works/{encoded}", opener=opener, max_bytes=1_000_000
    )
    try:
        payload = json.loads(raw.decode("utf-8"))
    except Exception as exc:
        raise PublicationError("Crossref DOI verification returned invalid JSON") from exc
    message = payload.get("message") or {}
    resolved = normalize_doi(message.get("DOI"))
    if resolved != norm:
        raise PublicationError("Crossref DOI verification did not resolve requested DOI")
    title = message.get("title") or []
    if isinstance(title, list):
        title = title[0] if title else ""
    return {
        "doi": norm,
        "status": "PASS",
        "resolved_title": " ".join(str(title).split())[:500],
        "route": "crossref-doi",
    }


def validate_publication_manifest(manifest: dict[str, Any], job_id: str) -> None:
    validate_job_id(job_id)
    if manifest.get("schema_version") != "cosci.job/1.0" or manifest.get("job_id") != job_id:
        raise PublicationError("private publication manifest identity mismatch")
    if manifest.get("privacy_class") != "PRIVATE_RESEARCH" or manifest.get("task_type") != "PUBLICATION":
        raise PublicationError("publication job must be private PUBLICATION")
    cost = manifest.get("cost_policy") or {}
    if cost.get("require_zero_cost") is not True or cost.get("automatic_paid_services") is not False:
        raise PublicationError("publication job violates zero-cost policy")
    params = manifest.get("parameters") or {}
    req = params.get("journal_requirements") or {}
    if not req.get("official_url"):
        raise PublicationError("publication job requires official journal requirements URL")
    artifacts = params.get("source_artifacts")
    if not isinstance(artifacts, list) or not artifacts:
        raise PublicationError("publication job requires source_artifacts")
    for item in artifacts:
        if (
            not isinstance(item, dict)
            or not item.get("drive_file_id")
            or not re.fullmatch(r"[0-9a-f]{64}", str(item.get("sha256") or ""))
        ):
            raise PublicationError("each source artifact requires Drive file ID and SHA-256")
        name = str(item.get("name") or "")
        if not _SAFE_NAME_RE.fullmatch(name):
            raise PublicationError("source artifact name must be a safe basename")
    ms = params.get("manuscript") or {}
    if not isinstance(ms.get("title"), str) or not ms.get("title", "").strip():
        raise PublicationError("canonical manuscript title is required")
    if not isinstance(ms.get("sections"), list) or not ms.get("sections"):
        raise PublicationError("canonical manuscript sections are required")
    fixture_status = "ENGINE_ACCEPTANCE_FIXTURE_NOT_SCIENTIFIC_MANUSCRIPT"
    if params.get("acceptance_fixture") is True and ms.get("document_status") != fixture_status:
        raise PublicationError(
            "acceptance fixture must carry the mandatory non-scientific-manuscript label"
        )
    if params.get("acceptance_fixture") is not True and ms.get("document_status") == fixture_status:
        raise PublicationError(
            "fixture document status cannot be used for a non-fixture publication job"
        )


def _download_source_artifacts(
    client: Any,
    specs: list[dict[str, Any]],
    root: Path,
) -> tuple[dict[str, Path], list[dict[str, Any]]]:
    paths: dict[str, Path] = {}
    refs: list[dict[str, Any]] = []
    source_dir = root / "source_artifacts"
    source_dir.mkdir(parents=True, exist_ok=True)
    for spec in specs:
        name = str(spec["name"])
        path = source_dir / name
        client.download(str(spec["drive_file_id"]), str(path))
        actual = sha256_file(path)
        expected = str(spec["sha256"]).lower()
        if actual.lower() != expected:
            raise PublicationError(f"source artifact SHA-256 mismatch for {name}")
        paths[name] = path
        refs.append(
            {
                "name": name,
                "kind": str(spec.get("kind") or "GENERIC"),
                "drive_file_id": str(spec["drive_file_id"]),
                "sha256": actual,
                "bytes": path.stat().st_size,
            }
        )
    return paths, refs


def build_canonical_source(
    params: dict[str, Any],
    source_refs: list[dict[str, Any]],
    requirements: dict[str, Any],
) -> dict[str, Any]:
    ms = params["manuscript"]
    source_names = {x["name"] for x in source_refs}
    citations: list[dict[str, Any]] = []
    seen_cits: set[str] = set()
    for cit in ms.get("citations") or []:
        cid = str(cit.get("citation_id") or "")
        if not cid or cid in seen_cits:
            raise PublicationError("citation IDs must be present and unique")
        seen_cits.add(cid)
        evidence_artifact = cit.get("evidence_artifact")
        if evidence_artifact and evidence_artifact not in source_names:
            raise PublicationError(f"citation {cid} references unknown evidence artifact")
        doi = normalize_doi(cit.get("doi")) if cit.get("doi") else None
        if cit.get("doi") and not doi:
            raise PublicationError(f"citation {cid} contains invalid DOI")
        citations.append(
            {
                "citation_id": cid,
                "doi": doi,
                "title": str(cit.get("title") or "").strip(),
                "evidence_artifact": evidence_artifact,
            }
        )
    figures = []
    for fig in ms.get("figures") or []:
        art = str(fig.get("artifact_name") or "")
        if art not in source_names:
            raise PublicationError(f"figure references unknown source artifact: {art}")
        figures.append(
            {
                "figure_id": str(fig.get("figure_id") or ""),
                "caption": str(fig.get("caption") or ""),
                "artifact_name": art,
            }
        )
    tables = []
    for table in ms.get("tables") or []:
        cols = [str(x) for x in table.get("columns") or []]
        rows = table.get("rows") or []
        if not cols or any(
            not isinstance(row, list) or len(row) != len(cols) for row in rows
        ):
            raise PublicationError("table rows must match declared columns")
        tables.append(
            {
                "table_id": str(table.get("table_id") or ""),
                "caption": str(table.get("caption") or ""),
                "columns": cols,
                "rows": [[str(v) for v in row] for row in rows],
                "result_artifact": table.get("result_artifact"),
            }
        )
    sections = []
    for section in ms["sections"]:
        paragraphs = []
        for para in section.get("paragraphs") or []:
            text = str(para.get("text") or "").strip()
            refs = [str(x) for x in para.get("citation_refs") or []]
            unknown = [x for x in refs if x not in seen_cits]
            if unknown:
                raise PublicationError(
                    f"paragraph references unknown citation IDs: {unknown}"
                )
            result_refs = [str(x) for x in para.get("result_refs") or []]
            unknown_results = [x for x in result_refs if x not in source_names]
            if unknown_results:
                raise PublicationError(
                    f"paragraph references unknown result artifacts: {unknown_results}"
                )
            paragraphs.append(
                {
                    "text": text,
                    "citation_refs": refs,
                    "result_refs": result_refs,
                }
            )
        sections.append(
            {
                "section_id": str(section.get("section_id") or ""),
                "heading": str(section.get("heading") or ""),
                "paragraphs": paragraphs,
            }
        )
    return {
        "schema_version": "cosci.manuscript-source/1.0",
        "document_status": str(ms.get("document_status") or "DRAFT_PRIVATE"),
        "title": str(ms["title"]).strip(),
        "authors": [str(x) for x in ms.get("authors") or []],
        "sections": sections,
        "citations": citations,
        "figures": figures,
        "tables": tables,
        "source_artifacts": source_refs,
        "journal_key": requirements["journal_key"],
        "journal_requirements_sha256": canonical_sha256(requirements),
        "journal_overrides": params.get("journal_overrides") or {},
        "availability_statement": str(ms.get("availability_statement") or ""),
    }


def build_citation_ledger(
    source: dict[str, Any],
    *,
    opener: Callable[..., Any] = urllib.request.urlopen,
) -> dict[str, Any]:
    ledger = []
    for cit in source["citations"]:
        if cit.get("doi"):
            verification = verify_doi(cit["doi"], opener=opener)
        else:
            verification = {"status": "NO_DOI", "doi": None}
        ledger.append({**cit, "verification": verification})
    failed = [
        x["citation_id"]
        for x in ledger
        if x["verification"]["status"] not in {"PASS", "NO_DOI"}
    ]
    if failed:
        raise PublicationError(
            f"citation verification failed for {len(failed)} citation(s)"
        )
    return {
        "schema_version": "cosci.citation-ledger/1.0",
        "citation_count": len(ledger),
        "verified_doi_count": sum(
            1 for x in ledger if x["verification"]["status"] == "PASS"
        ),
        "citations": ledger,
    }


def build_figure_table_registry(source: dict[str, Any]) -> dict[str, Any]:
    source_map = {x["name"]: x for x in source["source_artifacts"]}
    figures = []
    for fig in source["figures"]:
        art = source_map[fig["artifact_name"]]
        figures.append(
            {
                **fig,
                "source_sha256": art["sha256"],
                "source_drive_file_id": art["drive_file_id"],
                "integration": "EMBEDDED_AND_SEPARATE_COPY",
            }
        )
    tables = []
    for table in source["tables"]:
        result_name = table.get("result_artifact")
        if result_name and result_name not in source_map:
            raise PublicationError(
                "table result_artifact is not a canonical source artifact"
            )
        tables.append(
            {
                **table,
                "result_source": source_map.get(result_name) if result_name else None,
                "integration": "EMBEDDED_EDITABLE",
            }
        )
    return {
        "schema_version": "cosci.figure-table-registry/1.0",
        "figures": figures,
        "tables": tables,
    }


def _paragraph_text(para: dict[str, Any]) -> str:
    text = para["text"]
    refs = para.get("citation_refs") or []
    if refs:
        text += " [" + "; ".join(refs) + "]"
    return text


def _normalize_docx_zip(path: Path) -> None:
    temp = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(
        temp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as zout:
        for name in sorted(zin.namelist()):
            data = zin.read(name)
            info = zipfile.ZipInfo(
                filename=name, date_time=(1980, 1, 1, 0, 0, 0)
            )
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, data)
    os.replace(temp, path)


def render_docx(
    source: dict[str, Any],
    source_paths: dict[str, Path],
    out: Path,
) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except Exception as exc:
        raise PublicationError("python-docx rendering dependency unavailable") from exc
    doc = Document()
    props = doc.core_properties
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    props.created = fixed
    props.modified = fixed
    props.author = "Research CoScientist"
    props.last_modified_by = "Research CoScientist"
    props.title = source["title"]
    overrides = source.get("journal_overrides") or {}
    font_name = str(overrides.get("font_name") or "Times New Roman")
    font_size = float(overrides.get("font_size_pt") or 11)
    normal = doc.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(font_size)
    normal.paragraph_format.line_spacing = float(
        overrides.get("line_spacing") or 1.15
    )
    title = doc.add_heading(source["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if source.get("document_status"):
        paragraph = doc.add_paragraph(source["document_status"])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if source.get("authors"):
        paragraph = doc.add_paragraph(", ".join(source["authors"]))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    justify = (
        str(overrides.get("paragraph_alignment") or "LEFT").upper()
        == "JUSTIFIED"
    )
    for section in source["sections"]:
        doc.add_heading(section["heading"], level=1)
        for para in section["paragraphs"]:
            paragraph = doc.add_paragraph(_paragraph_text(para))
            if justify:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fig_width = float(overrides.get("figure_width_inches") or 5.5)
    for fig in source["figures"]:
        doc.add_picture(
            str(source_paths[fig["artifact_name"]]), width=Inches(fig_width)
        )
        paragraph = doc.add_paragraph(
            f"{fig['figure_id']}. {fig['caption']}"
        )
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tbl in source["tables"]:
        doc.add_paragraph(f"{tbl['table_id']}. {tbl['caption']}")
        table = doc.add_table(rows=1, cols=len(tbl["columns"]))
        table.style = "Table Grid"
        for idx, col in enumerate(tbl["columns"]):
            table.rows[0].cells[idx].text = col
        for row in tbl["rows"]:
            cells = table.add_row().cells
            for idx, value in enumerate(row):
                cells[idx].text = value
    if source["citations"]:
        doc.add_heading("References", level=1)
        for cit in source["citations"]:
            label = (
                f"[{cit['citation_id']}] "
                f"{cit.get('title') or 'Untitled source'}"
            )
            if cit.get("doi"):
                label += f". https://doi.org/{cit['doi']}"
            doc.add_paragraph(label)
    if source.get("availability_statement"):
        doc.add_heading("Data Availability", level=1)
        doc.add_paragraph(source["availability_statement"])
    doc.save(str(out))
    _normalize_docx_zip(out)


def render_pdf(
    source: dict[str, Any],
    source_paths: dict[str, Path],
    out: Path,
) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.pdfgen import canvas
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Image,
            Table,
            TableStyle,
        )
    except Exception as exc:
        raise PublicationError("ReportLab rendering dependency unavailable") from exc

    class InvariantCanvas(canvas.Canvas):
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            kwargs["invariant"] = 1
            super().__init__(*args, **kwargs)
            self.setAuthor("Research CoScientist")
            self.setTitle(source["title"])
            self.setCreator("Research CoScientist Publication Engine")

    overrides = source.get("journal_overrides") or {}
    styles = getSampleStyleSheet()
    body_align = (
        TA_JUSTIFY
        if str(overrides.get("paragraph_alignment") or "LEFT").upper()
        == "JUSTIFIED"
        else TA_LEFT
    )
    font_size = float(overrides.get("font_size_pt") or 10)
    body = ParagraphStyle(
        "BodyDet",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=font_size,
        leading=font_size * float(overrides.get("line_spacing") or 1.2),
        alignment=body_align,
    )
    heading = ParagraphStyle(
        "HeadingDet",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=17,
    )
    centered = ParagraphStyle("CenteredDet", parent=body, alignment=TA_CENTER)
    doc = SimpleDocTemplate(
        str(out),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
        title=source["title"],
        author="Research CoScientist",
    )
    story: list[Any] = [
        Paragraph(html.escape(source["title"]), styles["Title"]),
        Spacer(1, 8),
    ]
    if source.get("document_status"):
        story += [
            Paragraph(html.escape(source["document_status"]), centered),
            Spacer(1, 8),
        ]
    if source.get("authors"):
        story += [
            Paragraph(html.escape(", ".join(source["authors"])), centered),
            Spacer(1, 12),
        ]
    for section in source["sections"]:
        story.append(Paragraph(html.escape(section["heading"]), heading))
        for para in section["paragraphs"]:
            story += [
                Paragraph(html.escape(_paragraph_text(para)), body),
                Spacer(1, 6),
            ]
    fig_width = float(overrides.get("figure_width_inches") or 5.5) * inch
    for fig in source["figures"]:
        image = Image(str(source_paths[fig["artifact_name"]]))
        ratio = image.imageHeight / max(image.imageWidth, 1)
        image.drawWidth = fig_width
        image.drawHeight = fig_width * ratio
        story += [
            image,
            Paragraph(
                html.escape(f"{fig['figure_id']}. {fig['caption']}"), centered
            ),
            Spacer(1, 10),
        ]
    for tbl in source["tables"]:
        story += [
            Paragraph(html.escape(f"{tbl['table_id']}. {tbl['caption']}"), body)
        ]
        data = [tbl["columns"]] + tbl["rows"]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story += [table, Spacer(1, 10)]
    if source["citations"]:
        story.append(Paragraph("References", heading))
        for cit in source["citations"]:
            label = (
                f"[{cit['citation_id']}] "
                f"{cit.get('title') or 'Untitled source'}"
            )
            if cit.get("doi"):
                label += f". https://doi.org/{cit['doi']}"
            story += [Paragraph(html.escape(label), body), Spacer(1, 4)]
    if source.get("availability_statement"):
        story += [
            Paragraph("Data Availability", heading),
            Paragraph(html.escape(source["availability_statement"]), body),
        ]
    doc.build(story, canvasmaker=InvariantCanvas)


def _upload_once(client: Any, folder_id: str, path: Path, name: str) -> str:
    existing = _optional_child(client, folder_id, name)
    if existing:
        check = path.parent / f"drive-existing-{name}"
        client.download(existing["id"], str(check))
        if sha256_file(check) != sha256_file(path):
            raise PublicationError(f"existing private Drive artifact differs for {name}")
        return existing["id"]
    return client.upload(str(path), folder_id, name)


def run_private_publication_job(
    job_id: str,
    *,
    client: Any | None = None,
    workdir: str | Path | None = None,
    opener: Callable[..., Any] = urllib.request.urlopen,
) -> dict[str, Any]:
    validate_job_id(job_id)
    registry = load_publication_registry()
    client = client or _drive_client_from_env()
    job_folder = _search_job_folder(client, job_id)
    manifest_obj = _child(client, job_folder["id"], "job.json", folder=False)
    checkpoint_folder = _child(client, job_folder["id"], "checkpoints", folder=True)
    results_folder = _child(client, job_folder["id"], "results", folder=True)
    logs_folder = _child(client, job_folder["id"], "logs", folder=True)
    outputs_folder = _child(client, job_folder["id"], "outputs", folder=True)
    metadata_folder = _child(client, job_folder["id"], "metadata", folder=True)

    context = (
        tempfile.TemporaryDirectory(prefix="cosci-publication-")
        if workdir is None
        else None
    )
    root = Path(context.name if context else workdir)
    root.mkdir(parents=True, exist_ok=True)
    try:
        manifest, manifest_sha = _download_json(
            client, manifest_obj["id"], root / "job.json"
        )
        validate_publication_manifest(manifest, job_id)
        existing_result = _optional_child(client, results_folder["id"], "result.json")
        if existing_result:
            result, _ = _download_json(
                client, existing_result["id"], root / "result-existing.json"
            )
            if (
                result.get("status") != "SUCCEEDED"
                or (result.get("provenance") or {}).get("manifest_sha256")
                != manifest_sha
            ):
                raise PublicationError(
                    "existing publication result does not match this exact manifest"
                )
            return {"job_id": job_id, "status": "SUCCEEDED", "resumed": True}

        params = manifest["parameters"]
        source_paths, source_refs = _download_source_artifacts(
            client, params["source_artifacts"], root
        )
        requirements = retrieve_journal_requirements(
            params["journal_requirements"], opener=opener
        )
        canonical = build_canonical_source(params, source_refs, requirements)
        citations = build_citation_ledger(canonical, opener=opener)
        figure_tables = build_figure_table_registry(canonical)

        req_path = root / "journal_requirements.json"
        canonical_path = root / "canonical_manuscript.json"
        citation_path = root / "citation_ledger.json"
        ft_path = root / "figure_table_registry.json"
        _write_json(req_path, requirements)
        _write_json(canonical_path, canonical)
        _write_json(citation_path, citations)
        _write_json(ft_path, figure_tables)

        docx_path = root / "manuscript.docx"
        pdf_path = root / "manuscript.pdf"
        render_docx(canonical, source_paths, docx_path)
        render_pdf(canonical, source_paths, pdf_path)

        docx2 = root / "manuscript-rebuild.docx"
        pdf2 = root / "manuscript-rebuild.pdf"
        render_docx(canonical, source_paths, docx2)
        render_pdf(canonical, source_paths, pdf2)
        if (
            sha256_file(docx_path) != sha256_file(docx2)
            or sha256_file(pdf_path) != sha256_file(pdf2)
        ):
            raise PublicationError(
                "publication renderer failed deterministic rebuild check"
            )

        build_manifest = {
            "schema_version": "cosci.publication-build/1.0",
            "canonical_source_sha256": sha256_file(canonical_path),
            "journal_requirements_sha256": sha256_file(req_path),
            "citation_ledger_sha256": sha256_file(citation_path),
            "figure_table_registry_sha256": sha256_file(ft_path),
            "docx_sha256": sha256_file(docx_path),
            "pdf_sha256": sha256_file(pdf_path),
            "source_artifacts": source_refs,
            "rendering": registry["rendering"],
            "journal_overrides": canonical.get("journal_overrides") or {},
            "manubot_quarto_concepts": registry["manubot_quarto_concepts"],
            "deterministic_rebuild": "PASS",
            "privacy_class": "PRIVATE_RESEARCH",
        }
        build_path = root / "build_manifest.json"
        _write_json(build_path, build_manifest)

        checkpoint = {
            "schema_version": "cosci.checkpoint/1.0",
            "job_id": job_id,
            "privacy_class": "PRIVATE_RESEARCH",
            "stage": "PUBLICATION_PACKAGE_BUILT",
            "manifest_sha256": manifest_sha,
            "canonical_source_sha256": build_manifest[
                "canonical_source_sha256"
            ],
            "build_manifest_sha256": sha256_file(build_path),
            "resume_from": "ADVERSARIAL_REVIEW",
        }
        cp_path = root / "checkpoint.json"
        _write_json(cp_path, checkpoint)
        cp_id = _upload_once(
            client, checkpoint_folder["id"], cp_path, "checkpoint.json"
        )

        output_paths = [
            req_path,
            canonical_path,
            citation_path,
            ft_path,
            docx_path,
            pdf_path,
            build_path,
        ]
        output_refs = []
        for path in output_paths:
            output_refs.append(
                {
                    "name": path.name,
                    "drive_file_id": _upload_once(
                        client, outputs_folder["id"], path, path.name
                    ),
                    "sha256": sha256_file(path),
                    "bytes": path.stat().st_size,
                }
            )
        figure_copies = []
        for fig in canonical["figures"]:
            src = source_paths[fig["artifact_name"]]
            copy_name = f"submission-{fig['figure_id']}{src.suffix.lower()}"
            figure_copies.append(
                {
                    "figure_id": fig["figure_id"],
                    "name": copy_name,
                    "drive_file_id": _upload_once(
                        client, outputs_folder["id"], src, copy_name
                    ),
                    "sha256": sha256_file(src),
                }
            )

        provenance = {
            "schema_version": "cosci.publication-provenance/1.0",
            "manifest_sha256": manifest_sha,
            "publication_registry_sha256": sha256_file(REGISTRY_PATH),
            "canonical_source_sha256": build_manifest[
                "canonical_source_sha256"
            ],
            "build_manifest_sha256": sha256_file(build_path),
            "journal_requirements": {
                "official_url": requirements["official_url"],
                "content_sha256": requirements["content_sha256"],
                "marker_validation": requirements["marker_validation"],
            },
            "citation_verification_route": "Crossref keyless DOI API",
            "source_artifacts": source_refs,
            "outputs": output_refs,
            "figure_submission_copies": figure_copies,
            "executor": "GITHUB_ACTIONS_PUBLIC",
            "policy": {
                "zero_cost_only": True,
                "automatic_paid_routes": False,
                "private_manuscript_only": True,
                "final_scientific_judgment": "RESEARCH_DIRECTOR",
            },
        }
        prov_path = root / "publication_provenance.json"
        _write_json(prov_path, provenance)
        prov_id = _upload_once(
            client,
            metadata_folder["id"],
            prov_path,
            "publication_provenance.json",
        )

        result = {
            "schema_version": "cosci.result/1.0",
            "job_id": job_id,
            "project_id": manifest.get("project_id"),
            "status": "SUCCEEDED",
            "privacy_class": "PRIVATE_RESEARCH",
            "task_type": "PUBLICATION",
            "outputs": output_refs,
            "metadata": {
                "publication_provenance_file_id": prov_id,
                "figure_submission_copies": figure_copies,
            },
            "metrics": {
                "section_count": len(canonical["sections"]),
                "citation_count": citations["citation_count"],
                "verified_doi_count": citations["verified_doi_count"],
                "figure_count": len(canonical["figures"]),
                "table_count": len(canonical["tables"]),
            },
            "validation": {
                "passed": True,
                "checks": [
                    {"name": "zero_cost_policy", "status": "PASS"},
                    {
                        "name": "official_journal_requirements_retrieved",
                        "status": "PASS",
                    },
                    {
                        "name": "journal_requirements_provenance",
                        "status": "PASS",
                    },
                    {
                        "name": "source_artifact_hash_validation",
                        "status": "PASS",
                    },
                    {
                        "name": "single_canonical_manuscript_source",
                        "status": "PASS",
                    },
                    {
                        "name": "citation_identifier_verification",
                        "status": "PASS",
                    },
                    {
                        "name": "figure_table_traceability",
                        "status": "PASS",
                    },
                    {"name": "docx_render", "status": "PASS"},
                    {"name": "pdf_render", "status": "PASS"},
                    {"name": "deterministic_rebuild", "status": "PASS"},
                    {"name": "private_drive_publication", "status": "PASS"},
                ],
            },
            "provenance": {
                "manifest_sha256": manifest_sha,
                "canonical_source_sha256": build_manifest[
                    "canonical_source_sha256"
                ],
                "build_manifest_sha256": sha256_file(build_path),
                "executor": "GITHUB_ACTIONS_PUBLIC",
            },
            "next_state": "READY_FOR_ADVERSARIAL_REVIEW",
        }
        result_path = root / "result.json"
        _write_json(result_path, result)
        client.upload(str(result_path), results_folder["id"], "result.json")

        log_path = root / "execution.log"
        log_path.write_text(
            "\n".join(
                [
                    f"job_id={job_id}",
                    "task_type=PUBLICATION",
                    "official_requirements=PASS",
                    "source_hashes=PASS",
                    "canonical_source=PASS",
                    "citation_verification=PASS",
                    "figure_table_traceability=PASS",
                    "docx_render=PASS",
                    "pdf_render=PASS",
                    "deterministic_rebuild=PASS",
                    "private_drive_publication=PASS",
                    "next_state=READY_FOR_ADVERSARIAL_REVIEW",
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        client.upload(str(log_path), logs_folder["id"], "execution.log")
        return {
            "job_id": job_id,
            "status": "SUCCEEDED",
            "resumed": False,
            "checkpoint_file_id": cp_id,
            "result": result,
        }
    finally:
        if context:
            context.cleanup()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Research CoScientist private reproducible publication executor"
    )
    parser.add_argument("--job-id", required=True)
    parser.add_argument("--public-status", required=True)
    args = parser.parse_args(argv)
    run_private_publication_job(args.job_id)
    write_public_status(args.job_id, args.public_status, task_type="PUBLICATION")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
